"""WP3: server-side workspace field mappings, ATT&CK tag lint, query status lifecycle, deck slide 2 + branding,
research review states."""

import base64
import csv
import io
import os
import tempfile
from pathlib import Path

if "DATABASE_URL" not in os.environ:  # standalone run; under a full pytest run test_smoke already set a throwaway DB
    _tmp = Path(tempfile.mkdtemp())
    os.environ["DATABASE_URL"] = f"sqlite:///{(_tmp / 'test.db').as_posix()}"
os.environ["ANTHROPIC_API_KEY"] = ""

from fastapi.testclient import TestClient  # noqa: E402

from app import detection  # noqa: E402
from app.main import app  # noqa: E402

RID = "TR-2026-0142"
REVIEWER = {"X-User": "riyer"}
HUNTER = {"X-User": "averma"}


def client() -> TestClient:
    return TestClient(app)


# ------------------------------------------------------------------ B02 mapping (unit)

PROC = {"category": "process_creation", "conditions": [
    {"field": "parent_image", "op": "endswith", "values": ["\\w3wp.exe"]},
    {"field": "command_line", "op": "contains", "values": ["index=endpoint", "CommandLine"]}]}


def test_mapping_is_token_aware():
    spl = detection.translate(PROC, "spl", 30, "t")
    m = {"spl": {"index=endpoint": "index=acme_edr", "CommandLine": "cmdline",
                 'sourcetype="XmlWinEventLog:Microsoft-Windows-Sysmon/Operational"': "sourcetype=sysmon"}}
    out = detection.apply_mappings(spl, "spl", m)
    assert out.startswith("index=acme_edr sourcetype=sysmon EventCode=1")
    assert '"*index=endpoint*"' in out and '"*CommandLine*"' in out  # searched values are never rewritten
    assert "cmdline=" in out and "values(cmdline)" in out  # field renamed wherever it is a field
    # whole-value only; spacing / quoting variants of the same assignment match
    assert detection.apply_mappings("index=endpoint_old OR index = \"endpoint\"", "spl", {"spl": {"index=endpoint": "index=a"}}) \
        == "index=endpoint_old OR index=a"
    # table names and dotted fields; "*" applies to every platform
    kql = detection.translate(PROC, "kql_defender", 30, "t")
    out = detection.apply_mappings(kql, "kql_defender", {"*": {"ProcessCommandLine": "Cmd_s"}, "kql_defender": {"DeviceProcessEvents": "Proc_CL"}})
    assert out.startswith("Proc_CL\n") and "Cmd_s has_any" in out and "InitiatingProcessCommandLine" in out
    esql = detection.translate(PROC, "esql", 30, "t")
    out = detection.apply_mappings(esql, "esql", {"esql": {"logs-endpoint.events.process-*": "acme-edr-*", "process.name": "proc.name"}})
    assert out.startswith("FROM acme-edr-*") and "process.parent.name" in out and "proc.name" in out
    # other platforms' mappings do not leak
    assert detection.apply_mappings(kql, "kql_defender", {"spl": {"DeviceProcessEvents": "X"}}) == kql
    mq = detection.map_query({"platform": "kql_defender", "body": kql, "techniques": []}, {"kql_defender": {"DeviceProcessEvents": "index=x"}})
    assert mq["mapping_applied"] and any("index=" in i for i in mq["mapped_lint"])  # lint re-runs on the mapped text


# ------------------------------------------------------------------ B11 ATT&CK lint

def test_attack_tag_lint():
    good = detection.to_sigma("x", PROC, "DO-1", ["T1505.003"], "")
    assert detection.lint("sigma", good) == []
    bad = detection.to_sigma("x", PROC, "DO-1", ["T1505.003", "T9999"], "")
    assert any("T9999" in i for i in detection.lint("sigma", bad))
    assert any("tactic" in i for i in detection.lint("sigma", good.replace("attack.t1505.003", "attack.not_a_tactic")))
    assert detection.lint("sigma", good.replace("attack.t1505.003", "attack.initial_access")) == []
    spl = detection.translate(PROC, "spl", 30, "t")
    assert detection.lint("spl", spl, ["T1059.001"]) == []
    assert any("Malformed" in i for i in detection.lint("spl", spl, ["1059"]))
    assert any("Unknown ATT&CK technique T1999" in i for i in detection.lint("spl", spl, ["T1999"]))


# ------------------------------------------------------------------ B02 API + exports

def test_mapped_queries_api_and_exports():
    with client() as c:
        plain = c.get(f"/api/research/{RID}/queries").json()
        assert plain["total"] > 0 and "mapped_body" not in plain["items"][0]
        assert plain["statuses"] == detection.QUERY_STATUSES
        d = c.get(f"/api/research/{RID}/queries", params={"ws": "acme", "platform": "spl"}).json()
        assert d["workspace_id"] == "acme" and d["items"]
        spl_endpoint = [q for q in d["items"] if "index=endpoint" in q["body"]]
        assert spl_endpoint, "seed has SPL queries on index=endpoint"
        for q in spl_endpoint:
            assert q["mapping_applied"] and "index=acme_edr" in q["mapped_body"] and "index=endpoint" not in q["mapped_body"]
            assert isinstance(q["mapped_lint"], list)
        assert c.get(f"/api/research/{RID}/queries", params={"ws": "nope"}).status_code == 404

        # CSV export for Acme ships mapped bodies
        r = c.get(f"/api/research/{RID}/export/queries_csv", params={"ws": "acme"})
        assert r.status_code == 200
        rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
        spl_rows = [x for x in rows if x["platform"] == "spl"]
        assert spl_rows and not any("index=endpoint" in x["query"] for x in spl_rows)
        assert any(x["workspace_mapping"] == "acme" for x in spl_rows)
        # no workspace: unmapped
        rows = list(csv.DictReader(io.StringIO(c.get(f"/api/research/{RID}/export/queries_csv").content.decode("utf-8-sig"))))
        assert any("index=endpoint" in x["query"] for x in rows if x["platform"] == "spl")

        # HTML (the PDF's source) appendix is mapped too
        html = c.get(f"/api/research/{RID}/export/html", params={"ws": "acme", "inline": True}).text
        assert "index=acme_edr" in html and "Field mappings for Acme Bank applied" in html

        # library copy
        lib = c.get("/api/library/queries", params={"platform": "spl", "ws": "contoso", "page_size": 200}).json()
        hit = [q for q in lib["items"] if "index=endpoint" in q["body"]]
        assert hit and all("index=contoso_sysmon" in q["mapped_body"] for q in hit)
        det = c.get(f"/api/library/queries/{hit[0]['id']}", params={"ws": "contoso"}).json()
        assert "index=contoso_sysmon" in det["mapped_body"] and "mapped_lint" in det


# ------------------------------------------------------------------ B07 status lifecycle

def _spl_query(c):
    qs = c.get(f"/api/research/{RID}/queries").json()["items"]
    return next(q for q in qs if q["platform"] == "spl" and q["origin"] != "reference" and not q["lint"]
                and q["status"] == "syntax_checked")


def test_query_status_lifecycle():
    with client() as c:
        q = _spl_query(c)
        url = f"/api/research/{RID}/queries/{q['id']}"
        assert c.patch(url, json={"status": "approved"}).status_code == 422
        assert c.patch(url, json={"status": "reference"}).status_code == 409
        r = c.patch(url, json={"status": "reviewed"})
        assert r.status_code == 200 and r.json()["status"] == "reviewed"
        assert c.get(f"/api/library/queries/{q['id']}").json()["status"] == "reviewed"
        # a body that fails lint cannot be marked reviewed, and editing resets the status
        assert c.patch(url, json={"body": q["body"] + " | project x", "status": "reviewed"}).status_code == 409

        # library: deploy it, then an unrelated record edit re-syncs; the library keeps the more advanced status
        assert c.patch(f"/api/library/queries/{q['id']}", json={"status": "bogus"}).status_code == 422
        r = c.patch(f"/api/library/queries/{q['id']}", json={"deployed_workspaces": ["acme"]})
        assert r.status_code == 200 and r.json()["status"] == "deployed"
        assert c.patch(f"/api/research/{RID}/record", json={"changes": {"tags": ["wp3"]}}).status_code == 200
        assert c.get(f"/api/library/queries/{q['id']}").json()["status"] == "deployed"
        # an explicit research-side change is authoritative (demotion propagates)
        assert c.patch(url, json={"status": "lab_tested"}).json()["status"] == "lab_tested"
        assert c.get(f"/api/library/queries/{q['id']}").json()["status"] == "lab_tested"

        # vendor reference queries
        ref = next((x for x in c.get(f"/api/research/{RID}/queries").json()["items"] if x["origin"] == "reference"), None)
        if ref:
            rurl = f"/api/research/{RID}/queries/{ref['id']}"
            assert c.patch(rurl, json={"status": "syntax_checked"}).status_code == 409
            assert c.patch(rurl, json={"status": "reviewed"}).status_code == 200


def test_ioc_verdict_override_survives_sync():
    with client() as c:
        items = c.get("/api/library/iocs", params={"q": "104.238.159"}).json()["items"]
        assert items
        iid = items[0]["id"]
        r = c.patch(f"/api/library/iocs/{iid}", json={"verdict": "benign"})
        assert r.status_code == 200 and r.json()["verdict"] == "benign" and r.json()["verdict_override"]["verdict"] == "benign"
        assert c.patch(f"/api/research/{RID}/record", json={"changes": {"tags": ["wp3", "sync"]}}).status_code == 200
        assert c.get(f"/api/library/iocs/{iid}").json()["verdict"] == "benign"
        r = c.patch(f"/api/library/iocs/{iid}", json={"clear_override": True})
        assert r.status_code == 200 and r.json()["verdict_override"] is None
        c.patch(f"/api/research/{RID}/record", json={"changes": {"tags": ["wp3"]}})
        assert c.get(f"/api/library/iocs/{iid}").json()["verdict"] != "benign"


# ------------------------------------------------------------------ B17 deck

PNG_1PX = base64.b64encode(bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360f8cfc0f01f0005000201e2e8a0b30000000049454e44ae426082")).decode()


def _deck_text(data: bytes) -> tuple[str, object]:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    return text, prs


def test_deck_queries_run_and_branding():
    from app.db import SessionLocal
    from app.models import Workspace

    with SessionLocal() as db:
        w = db.get(Workspace, "contoso")
        w.branding = {**(w.branding or {}), "logo_data_uri": f"data:image/png;base64,{PNG_1PX}"}
        db.commit()
    with client() as c:
        # Contoso recorded queries_run=["DO-3"]: slide 2 counts what was run
        r = c.get(f"/api/research/{RID}/export/pptx", params={"ws": "contoso"})
        assert r.status_code == 200
        text, prs = _deck_text(r.content)
        assert "Queries run by type" in text
        master = prs.slide_master
        assert any(sh.shape_type == 13 for sh in master.shapes), "logo picture on the master"
        assert any(sh.name == "ThreatLens brand band" and str(sh.fill.fore_color.rgb) == "B8790F" for sh in master.shapes)
        # Acme recorded nothing: fall back to generated counts, labelled
        text, _ = _deck_text(c.get(f"/api/research/{RID}/export/pptx", params={"ws": "acme"}).content)
        assert "Queries generated by type (none recorded as run)" in text


def test_queries_run_by_type_resolution():
    from app.exports.context import queries_run_by_type
    rec = {"hunts": {"queries": [{"id": "Q-1", "type": "ioc", "group": "IOC-IPV4"}, {"id": "Q-2", "type": "ioa", "group": "DET-1",
                                                                                     "opportunity_id": "DO-1"}]},
           "detection_opportunities": [{"id": "DO-1", "type": "ioa"}]}
    c = queries_run_by_type(rec, ["Q-1", "DO-1", "IOC-IPV4", "whatever"])
    assert c["ioc"] == 2 and c["ioa"] == 1 and c["other"] == 1


# ------------------------------------------------------------------ B05 review states

def test_review_states_and_publish():
    with client() as c:
        draft = next(x for x in c.get("/api/research", params={"status": ["draft"]}).json()["items"])
        rid = draft["id"]
        r = c.post(f"/api/research/{rid}/status", json={"action": "publish"}, headers=REVIEWER)
        assert r.status_code == 409 and "in review" in r.json()["detail"]
        assert c.post(f"/api/research/{rid}/review", json={"section": "executive_summary"}, headers=HUNTER).status_code == 403
        assert c.post(f"/api/research/{rid}/review", json={"section": "executive_summary", "state": "nope"}, headers=REVIEWER).status_code == 422
        assert c.post(f"/api/research/{rid}/review", json={"section": "executive_summary"}, headers=REVIEWER).status_code == 200
        assert c.post(f"/api/research/{rid}/status", json={"action": "submit"}, headers=HUNTER).status_code == 200
        # Grounding gate (test_r1): sample drafts cite no sources, so the reviewer explicitly approves the blocking sections.
        ready = c.get(f"/api/research/{rid}/readiness").json()
        if ready["blocking"]:
            assert c.post(f"/api/research/{rid}/status", json={"action": "publish"}, headers=REVIEWER).status_code == 409
        for section in {i["section"] for i in ready["issues"] if i["severity"] == "block"}:
            assert c.post(f"/api/research/{rid}/review", json={"section": section}, headers=REVIEWER).status_code == 200
        before =c.get(f"/api/research/{rid}/versions").json()
        r = c.post(f"/api/research/{rid}/status", json={"action": "publish", "note": "LGTM"}, headers=REVIEWER)
        assert r.status_code == 200 and r.json()["status"] == "published"
        after = c.get(f"/api/research/{rid}/versions").json()
        assert len(after) == len(before) + 1 and after[0]["summary"].startswith("Published by")
        assert after[0]["changed_by"]["id"] == "riyer"
        d = c.get(f"/api/research/{rid}").json()
        assert d["status"] == "published" and d["record"]["status"] == "published"
