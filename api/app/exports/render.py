"""Export renderers (spec section 10): HTML report, PDF, email HTML/.eml, 2-slide PPTX, JSON, CSV, period XLSX + summary PDF."""

from __future__ import annotations

import base64
import binascii
import csv
import io
import json
import re
from email.message import EmailMessage
from html import escape as html_escape
from email.policy import SMTP
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape

from ..sources import run_guarded_page
from . import context as C

_env = Environment(loader=FileSystemLoader(Path(__file__).parent / "templates"), autoescape=select_autoescape(["html"]))


class ExportRefused(Exception):
    pass


def report_html(vm: dict) -> str:
    return _env.get_template("report.html").render(**vm)


def email_html(vm: dict) -> str:
    if vm["tlp"] == "RED":
        raise ExportRefused("TLP:RED reports cannot be emailed")
    return _env.get_template("email.html").render(**vm, logo_data_uri=vm["branding"].get("logo_data_uri"))


def eml(vm: dict) -> bytes:
    msg = EmailMessage(policy=SMTP)
    msg["Subject"] = f"[TLP:{vm['tlp']}] {vm['rec'].get('title', vm['r'].id)}"
    msg["X-Unsent"] = "1"  # Outlook opens it as a draft
    msg["To"] = ""
    msg.set_content(f"{vm['rec'].get('title')}\n\n{vm['rec'].get('executive_summary', '')}\n\nFull report: {vm['report_url']}")
    msg.add_alternative(email_html(vm), subtype="html")
    return bytes(msg)


def html_to_pdf(html: str, header_left: str, header_right: str, footer_left: str) -> bytes:
    header_left, header_right, footer_left = (html_escape(x) for x in (header_left, header_right, footer_left))
    style = "font-family:Segoe UI,Arial,sans-serif;font-size:7.5pt;color:#525B6B;width:100%;padding:0 16mm;display:flex;justify-content:space-between;"
    header = f'<div style="{style}"><span>{header_left}</span><span style="font-weight:600;color:#2A303B">{header_right}</span></div>'
    footer = (f'<div style="{style}"><span>{footer_left}</span><span>Page <span class="pageNumber"></span> of '
              f'<span class="totalPages"></span></span><span style="font-weight:600;color:#2A303B">{header_right}</span></div>')
    async def work(page):
        # The report is built from research content (LLM output, article text, hunter edits): it may load only inline
        # data and public http(s) resources (fonts), fetched through the pinned SSRF-checked client.
        await page.set_content(html, wait_until="networkidle", timeout=30000)
        return await page.pdf(format="A4", print_background=True, display_header_footer=True, header_template=header,
                              footer_template=footer, margin={"top": "18mm", "bottom": "18mm", "left": "16mm", "right": "16mm"})

    return run_guarded_page(work, timeout=30)


def pdf(vm: dict) -> bytes:
    tlp = f"TLP:{vm['tlp']}"
    return html_to_pdf(report_html(vm), vm["client"], tlp, f"{vm['r'].id} · {vm['client']}")


def record_json(vm: dict) -> bytes:
    rec = {k: v for k, v in vm["rec"].items() if not k.startswith("_")}
    return json.dumps(rec, indent=2, ensure_ascii=False, default=str).encode()


def iocs_csv(vm: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["research_id", "type", "value_defanged", "verdict", "role", "context", "sources", "reputation"])
    for i in vm["rec"].get("iocs", []):
        w.writerow([vm["r"].id, i["type"], i["value"], i.get("verdict", ""), i.get("role", ""), i.get("context", ""),
                    " ".join(i.get("source_ids", [])), i.get("reputation_summary", "")])
    return buf.getvalue().encode("utf-8-sig")


def queries_csv(vm: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    # `query` is the copy-paste-ready text: the chosen workspace's field mappings are applied server-side and `lint` is
    # re-run on it. `workspace_mapping` names the workspace when its mappings changed the query.
    w.writerow(["research_id", "query_id", "group", "type", "platform", "title", "status", "provenance", "source_ids", "techniques",
                "log_sources", "fp_notes", "workspace_mapping", "lint", "query"])
    ws = vm.get("ws")
    for q in vm.get("hunt_queries", vm["rec"].get("hunts", {}).get("queries", [])):
        w.writerow([vm["r"].id, q["id"], q.get("group", ""), q["type"], q["platform"], q["title"], q.get("status", ""),
                    q.get("provenance", ""), " ".join(q.get("source_ids", [])), " ".join(q.get("techniques", [])),
                    "; ".join(q.get("log_sources", [])), q.get("fp_notes", ""), ws.id if ws and q.get("mapping_applied") else "",
                    "; ".join(q.get("lint") or []), q["body"]])
    return buf.getvalue().encode("utf-8-sig")


# ---------------------------------------------------------------- 2-slide deck

def _logo_stream(data_uri: str | None):
    """BytesIO of a raster logo from a workspace `logo_data_uri` (PNG/JPEG/GIF/BMP), else None."""
    if not data_uri:
        return None
    m = re.match(r"data:image/(png|jpe?g|gif|bmp|x-png);base64,(.+)$", data_uri.strip(), re.S | re.I)
    if not m:
        return None
    try:
        return io.BytesIO(base64.b64decode(m.group(2)))
    except (ValueError, binascii.Error):
        return None


def pptx(vm: dict) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    def rgb(h: str) -> RGBColor:
        return RGBColor.from_string(h.lstrip("#").upper())

    rec, r = vm["rec"], vm["r"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    FONT = "Atkinson Hyperlegible Next"
    brand = vm.get("brand_color") or "#5249A8"
    if not re.fullmatch(r"#[0-9A-Fa-f]{6}", brand):
        brand = "#5249A8"
    logo = _logo_stream(vm.get("branding", {}).get("logo_data_uri"))

    # Client-brandable master (spec 10): the workspace colour band and logo live on the slide master itself, so every
    # slide (and any slide the client adds from the Blank layout) carries them.
    from pptx.shapes.shapetree import SlideShapes

    sm = SlideShapes(prs.slide_master._element.cSld.spTree, prs.slide_master)  # MasterShapes has no add_* methods
    band = sm.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.06))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string(brand.lstrip("#").upper())
    band.line.fill.background()
    band.name = "ThreatLens brand band"
    logo_w = 0.0
    if logo is not None:
        try:
            pic = sm.add_picture(logo, Inches(0.5), Inches(7.02), height=Inches(0.34))
            if pic.width > Inches(1.6):  # keep very wide logos inside the footer
                ratio = Inches(1.6) / pic.width
                pic.width, pic.height = Inches(1.6), int(pic.height * ratio)
            pic.name = "Client logo"
            logo_w = pic.width / 914400 + 0.15
        except Exception:  # noqa: BLE001 - unsupported image type (e.g. SVG): deck still renders without a logo
            logo_w = 0.0

    def text(slide, x, y, w, h, s, size=12, bold=False, color="#181C23", mono=False, align=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = s if isinstance(s, list) else [s]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = "Atkinson Hyperlegible Mono" if mono else FONT
            run.font.color.rgb = rgb(color)
            if align:
                p.alignment = align
        return tb

    def badge(slide, x, y, label, soft, fg, square=None, w=1.4):
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.32))
        shp.adjustments[0] = 0.12
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(soft)
        shp.line.fill.background()
        tf = shp.text_frame
        tf.margin_left = Inches(0.25 if square else 0.08)
        run = tf.paragraphs[0].add_run()
        run.text = label
        run.font.size, run.font.bold, run.font.name = Pt(12), True, FONT
        run.font.color.rgb = rgb(fg)
        if square:
            sq = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.09), Inches(y + 0.11), Inches(0.1), Inches(0.1))
            sq.fill.solid()
            sq.fill.fore_color.rgb = rgb(square)
            sq.line.color.rgb = rgb("#8A93A3") if square.upper() == "#FFFFFF" else rgb(square)
        return x + w + 0.12

    def master(slide):
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.08), prs.slide_height)
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = rgb(vm["sev"]["solid"])
        stripe.line.fill.background()
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(6.95), Inches(12.83), Inches(6.95))
        line.line.color.rgb = rgb(brand)
        text(slide, 0.5 + logo_w, 7.0, 9 - logo_w, 0.35, f"{vm['client']}  ·  {r.id}  ·  {vm['now'][:10]}", 12, color="#6A7384")
        text(slide, 10.3, 7.0, 2.53, 0.35, f"TLP:{vm['tlp']}", 12, True, "#2A303B", align=PP_ALIGN.RIGHT)

    def rail(slide, y):
        seg_w = (12.33 - 13 * 0.04) / 14
        for i, t in enumerate(vm["rail"]):
            x = 0.5 + i * (seg_w + 0.04)
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(seg_w), Inches(0.12))
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(C.rail_fill(t["count"]))
            s.line.fill.background()
            text(slide, x - 0.05, y + 0.13, seg_w + 0.1, 0.3, t["short"], 12 if seg_w > 0.9 else 10, color="#6A7384", align=PP_ALIGN.CENTER)

    # Slide 1 — Threat overview
    s1 = prs.slides.add_slide(blank)
    master(s1)
    text(s1, 0.5, 0.4, 6.1, 1.2, rec.get("title", ""), 28, True)
    x = 0.5
    x = badge(s1, x, 1.75, vm["sev"]["label"], vm["sev"]["soft"], vm["sev"]["text"], vm["sev"]["solid"])
    x = badge(s1, x, 1.75, f"TLP:{vm['tlp']}", "#F0F2F5", "#2A303B", C.TLP_SQUARE.get(vm["tlp"], "#B8870B"), 1.7)
    badge(s1, x, 1.75, f"{(rec.get('confidence') or 'moderate').capitalize()} confidence", "#F0F2F5", "#2A303B", None, 2.1)
    summary = " ".join((rec.get("executive_summary") or "").split())
    words = summary.split()
    short = " ".join(words[:70]) + ("…" if len(words) > 70 else "")
    text(s1, 0.5, 2.3, 6.1, 2.9, short, 13, color="#3D4452")
    text(s1, 0.5, 5.3, 6.1, 0.3, "Impacted industries", 14, True)
    text(s1, 0.5, 5.62, 6.1, 0.5, vm["industries"], 12, color="#3D4452")

    text(s1, 6.9, 0.4, 5.9, 0.35, "Attack flow", 14, True)
    steps = []
    for p in rec.get("attack_paths", [])[:1]:
        steps = p["steps"][:6]
    if not steps:
        steps = [{"behaviour": m["procedure"] or m["technique"], "technique_id": m["technique_id"]} for m in rec.get("mitre", [])[:5]]
    box_h, gap = 0.52, 0.16
    for i, st in enumerate(steps):
        y = 0.85 + i * (box_h + gap)
        b = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(y), Inches(5.93), Inches(box_h))
        b.adjustments[0] = 0.15
        b.fill.solid()
        b.fill.fore_color.rgb = rgb("#F7F8FA")
        b.line.color.rgb = rgb("#CDD2DA")
        tf = b.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.12)
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{st['technique_id']}  "
        r1.font.size, r1.font.name, r1.font.bold = Pt(12), "Atkinson Hyperlegible Mono", True
        r1.font.color.rgb = rgb("#433B8E")
        r2 = p.add_run()
        beh = st["behaviour"]
        r2.text = beh if len(beh) < 70 else beh[:68] + "…"
        r2.font.size, r2.font.name = Pt(12), FONT
        r2.font.color.rgb = rgb("#181C23")
        if i < len(steps) - 1:
            c = s1.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(9.86), Inches(y + box_h), Inches(9.86), Inches(y + box_h + gap))
            c.line.color.rgb = rgb("#AEB5C1")
    vy = 0.85 + len(steps) * (box_h + gap) + 0.05
    cves = rec.get("vulnerabilities", [])[:4]
    if cves and vy < 5.4:
        text(s1, 6.9, vy, 5.9, 0.3, "Vulnerabilities", 14, True)
        text(s1, 6.9, vy + 0.32, 5.9, 1.2, [f"{v['cve']}  CVSS {v.get('cvss') or '—'}" for v in cves], 12, mono=True, color="#3D4452")
    rail(s1, 6.25)

    # Slide 2 — Hunt outcome and actions
    s2 = prs.slides.add_slide(blank)
    master(s2)
    text(s2, 0.5, 0.4, 12, 0.6, "Hunt outcome and actions", 28, True)
    res = vm["result"] or {}
    st = vm["result_style"]
    text(s2, 0.5, 1.2, 5, 0.3, f"Result for {vm['client']}", 14, True)
    b = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.6), Inches(3.2), Inches(0.55))
    b.adjustments[0] = 0.5
    b.fill.solid()
    b.fill.fore_color.rgb = rgb(st["soft"])
    b.line.fill.background()
    run = b.text_frame.paragraphs[0].add_run()
    run.text = st["label"]
    run.font.size, run.font.bold, run.font.name = Pt(18), True, FONT
    run.font.color.rgb = rgb(st["text"])
    b.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s2, 0.5, 2.3, 4.9, 1.2, res.get("summary") or "Hunt not yet recorded for this workspace.", 13, color="#3D4452")
    text(s2, 0.5, 3.6, 4.9, 0.3, "Hunt window", 14, True)
    text(s2, 0.5, 3.92, 4.9, 0.4, res.get("hunt_window") or "—", 12, color="#3D4452")
    # "# queries run by type" from the workspace result; generated counts (labelled) when no run was recorded.
    qb = vm.get("q_run_by_type") or vm["q_by_type"]
    text(s2, 0.5, 4.45, 4.9, 0.3, vm.get("q_run_label", "Queries by type"), 14, True)
    lines = [f"IoA  {qb.get('ioa', 0)}", f"IoC  {qb.get('ioc', 0)}", f"Vulnerability  {qb.get('vuln', 0)}", f"TTP  {qb.get('ttp', 0)}"]
    if qb.get("other"):
        lines.append(f"Other  {qb['other']}")
    text(s2, 0.5, 4.77, 4.9, 1.2, lines, 12, color="#3D4452")

    text(s2, 5.9, 1.2, 6.9, 0.3, "Top recommendations", 14, True)
    y = 1.6
    shown = 0
    for label, items in vm["grouped_recs"]:
        if not items or shown >= 5:
            continue
        text(s2, 5.9, y, 6.9, 0.3, label, 12, True, "#525B6B")
        y += 0.32
        for x_ in items[: 5 - shown]:
            t = text(s2, 5.9, y, 6.9, 0.6, f"•  {x_['action']}", 12, color="#181C23")
            y += 0.34 * max(1, (len(x_["action"]) // 85) + 1) + 0.04
            shown += 1
    y = max(y + 0.1, 4.6)
    n_iocs = len(rec.get("iocs", []))
    text(s2, 5.9, y, 3.3, 0.3, "Key indicators", 14, True)
    text(s2, 5.9, y + 0.32, 3.3, 0.4, f"{n_iocs} IoCs · {len(vm['key_iocs'])} malicious or suspicious", 12, color="#3D4452")
    text(s2, 9.4, y, 3.4, 0.3, "Log-source gaps", 14, True)
    gaps = sorted({g["data_source"].replace("_", " ") for g in vm["gaps"]})
    text(s2, 9.4, y + 0.32, 3.4, 0.9, ", ".join(gaps) or "None for this workspace", 12, color="#8F5A00" if gaps else "#3D4452")
    text(s2, 0.5, 6.35, 12.3, 0.4,
         f"Analyst: {vm['author'].name if vm['author'] else '—'}   ·   Reviewer: {vm['reviewer'].name if vm['reviewer'] else '—'}   ·   Full report: {vm['report_url']}",
         12, color="#6A7384")

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- period export

def period_xlsx(rows: list[dict], ttp_rows: list[dict], ioc_rows: list[dict], query_rows: list[dict], meta: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="F7F8FA")
    head_font = Font(bold=True, color="525B6B")

    def sheet(ws, headers, data):
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
        for row in data:
            ws.append([row.get(h.lower().replace(" ", "_"), "") for h in headers])
        for i, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(r.get(h.lower().replace(" ", "_"), ""))) for r in data[:200]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(60, width)
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=len(str(c.value or "")) > 60)

    ws = wb.active
    ws.title = "Runs"
    sheet(ws, ["ID", "Title", "Status", "Created", "Published", "Workspaces", "Severity", "Classification", "Actors", "CVEs", "Results", "Analyst"], rows)
    sheet(wb.create_sheet("TTPs"), ["Research ID", "Tactic", "Technique ID", "Technique", "Confidence"], ttp_rows)
    sheet(wb.create_sheet("IoCs"), ["Research ID", "Type", "Value", "Verdict", "Sources"], ioc_rows)
    sheet(wb.create_sheet("Queries"), ["Research ID", "Query ID", "Type", "Platform", "Title", "Status"], query_rows)
    info = wb.create_sheet("About")
    for k, v in meta.items():
        info.append([k, str(v)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def period_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    if rows:
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def period_summary_pdf(vm: dict) -> bytes:
    html = _env.get_template("period.html").render(**vm)
    return html_to_pdf(html, vm["scope"], "TLP:AMBER", f"Period summary · {vm['label']}")
